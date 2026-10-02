"""Build the updated user manual (manual.docx) for the Pilatus Integration GUI."""
import os
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
D = os.path.join(HERE, "build")                       # screenshots in, manual.docx out
GUI = os.path.dirname(os.path.dirname(HERE))           # repo root (icon.png)
OUT = os.path.join(D, "manual.docx")
RED = RGBColor(0xC8, 0x10, 0x2E)
GREY = "D9E2F3"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, m, Inches(1))
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
st.paragraph_format.space_after = Pt(6)
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
    doc.styles[name].font.name = "Calibri Light"
    doc.styles[name].font.size = Pt(size)


# ------------------------------------------------------------------ helpers
def p(text="", style=None, bold=False, italic=False, size=None, align=None, after=None, keep=False):
    para = doc.add_paragraph(style=style)
    if text:
        add_runs(para, text, bold=bold, italic=italic, size=size)
    if align:
        para.alignment = align
    if after is not None:
        para.paragraph_format.space_after = Pt(after)
    if keep:
        para.paragraph_format.keep_with_next = True
    return para


def add_runs(para, text, bold=False, italic=False, size=None):
    """**bold**, *italic*, and `code` inline markup."""
    import re
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**"):
            r = para.add_run(tok[2:-2]); r.bold = True
        elif tok.startswith("`"):
            r = para.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(10)
        elif tok.startswith("*"):
            r = para.add_run(tok[1:-1]); r.italic = True
        else:
            r = para.add_run(tok)
        if bold:
            r.bold = True
        if italic:
            r.italic = True
        if size:
            r.font.size = Pt(size)
    return para


def h(text, level=1, new_page=False):
    hd = doc.add_heading(text, level=level)
    if new_page:
        hd.paragraph_format.page_break_before = True
    return hd


def bullets(items, style="List Bullet"):
    for it in items:
        p(it, style=style, after=2)


def _new_number_list():
    """A fresh numbering instance (restarting at 1) on the 'List Number' definition."""
    numbering = doc.part.numbering_part.element
    style_numpr = doc.styles["List Number"].element.pPr.numPr
    base_num = style_numpr.numId.val
    abstract = None
    for num in numbering.findall(qn("w:num")):
        if num.get(qn("w:numId")) == str(base_num):
            abstract = num.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_id = max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(new_id))
    a = OxmlElement("w:abstractNumId"); a.set(qn("w:val"), abstract); num.append(a)
    ov = OxmlElement("w:lvlOverride"); ov.set(qn("w:ilvl"), "0")
    so = OxmlElement("w:startOverride"); so.set(qn("w:val"), "1"); ov.append(so); num.append(ov)
    numbering.append(num)
    return new_id


def steps(items):
    nid = _new_number_list()
    for it in items:
        para = p(it, style="List Number", after=2)
        pPr = para._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl"); il.set(qn("w:val"), "0"); numPr.append(il)
        ni = OxmlElement("w:numId"); ni.set(qn("w:val"), str(nid)); numPr.append(ni)
        pPr.append(numPr)


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def keep_together(t):
    """Keep a short table on one page: rows unsplit, each row kept with the next."""
    for i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
        if i < len(t.rows) - 1:
            for c in row.cells:
                for pp in c.paragraphs:
                    pp.paragraph_format.keep_with_next = True


def table(rows, widths, header=None, number_col=False, size=None, keep=None):
    t = doc.add_table(rows=0, cols=len(widths))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if header:
        cells = t.add_row().cells
        for c, txt in zip(cells, header):
            c.text = ""
            add_runs(c.paragraphs[0], txt, bold=True)
            shade(c, GREY)
    for r in rows:
        cells = t.add_row().cells
        for j, (c, txt) in enumerate(zip(cells, r)):
            c.text = ""
            para = c.paragraphs[0]
            if number_col and j == 0:
                run = para.add_run(txt); run.bold = True; run.font.color.rgb = RED; run.font.size = Pt(12)
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                parts = txt.split("\n")
                add_runs(para, parts[0], size=size)
                for extra in parts[1:]:
                    add_runs(c.add_paragraph(), extra, size=size)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(2)
    for row in t.rows:
        for c, wd in zip(row.cells, widths):
            c.width = Inches(wd)
    if header:                                   # repeat header row on page breaks
        trPr = t.rows[0]._tr.get_or_add_trPr()
        el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)
    if keep if keep is not None else (number_col or len(rows) <= 8):   # short tables stay in one piece
        keep_together(t)
    p(after=4)
    return t


def figure(img, caption, width=6.5):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(os.path.join(D, img), width=Inches(width))
    cap = p(caption, italic=True, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    return cap


def note(text, label="Note"):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    c.text = ""
    add_runs(c.paragraphs[0], f"**{label}.** " + text)
    shade(c, "F2F2F2")
    c.width = Inches(6.5)
    p(after=4)


def toc():
    para = doc.add_paragraph()
    for kind, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None),
                       (None, None), ("end", None)):
        if kind:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind)
            r = OxmlElement("w:r"); r.append(fc); para._p.append(r)
        elif text:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = text
            r = OxmlElement("w:r"); r.append(it); para._p.append(r)
        else:
            r = para.add_run("Right-click and choose Update Field to build the table of contents.")
            r.italic = True


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ title
title = doc.add_table(rows=1, cols=2)
c0, c1 = title.rows[0].cells
c0.width, c1.width = Inches(1.3), Inches(5.2)
c0.paragraphs[0].add_run().add_picture(os.path.join(GUI, "icon.png"), width=Inches(1.15))
r = c1.paragraphs[0].add_run("Pilatus Integration GUI"); r.font.size = Pt(28); r.font.name = "Calibri Light"
r = c1.add_paragraph().add_run("User Manual"); r.font.size = Pt(16); r.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
r = c1.add_paragraph().add_run("SSRL Beamline 2-1 · Pilatus 100K · updated October 2026")
r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
p(after=6)
p("Contents", bold=True, size=13, after=2)
toc()

# ------------------------------------------------------------------ 1 overview
h("1. Overview", new_page=True)
p("The Pilatus Integration GUI turns the Pilatus 100K images collected at SSRL beamline 2-1 into "
  "one-dimensional powder diffraction patterns (intensity vs 2θ) and displays them. The detector is "
  "scanned in 2θ during each measurement; the program combines every image of a scan into one "
  "pattern, normalized to the incident beam monitor, and writes it as an `.xye` file.")
p("With it you can:")
bullets([
    "integrate a single scan or a range of scans with one click;",
    "display patterns one at a time, overlaid, or as a contour (waterfall) plot;",
    "import patterns integrated earlier;",
    "run **live mode** during an experiment: each scan is integrated automatically as soon as it is "
    "complete and added to a live waterfall plot, with no clicks needed;",
    "**detect events**: flag scans where the pattern changed (a reaction, a phase transition, peaks "
    "sharpening or broadening), during a live run or afterwards on loaded data.",
])

# ------------------------------------------------------------------ 2 before you start
h("2. Before you start")
p("You need four things. The first three come from the beamline; the fourth is wherever you want the "
  "integrated data to go.")
table([
    ["Calibration file (`.cal`)", "Direct-beam position on the detector and the sample-to-detector distance. "
     "Made with the beamline calibration script from a direct-beam scan; beamline staff will tell you "
     "which file applies to your measurements. See Section 10."],
    ["SPEC file", "The data file written by SPEC during your measurements (no file extension). It "
     "holds the 2θ position and monitor counts for every image, and your user name."],
    ["Image folder", "The folder holding the Pilatus `.raw` images (and their `.pdi` files) for the "
     "scans in the SPEC file."],
    ["Output folder", "Where the integrated `.xye` files are written. Defaults to the folder of the "
     "SPEC file."],
], widths=[1.7, 4.8], header=["Item", "What it is"])
h("Starting the program", 2)
p("On the beamline computer, double-click `Pilatus_Integration_GUI.exe` (or its shortcut). Keep "
  "`manual.pdf` in the same folder as the program so that **Help > Open Manual PDF** can find it.")
p("To run from the Python source instead, open a terminal and run:")
p("`conda activate PilatusGUI`", after=0)
p("`python Pilatus_Integration_GUI.py`  (in the program folder)")

# ------------------------------------------------------------------ 3 layout
h("3. Program layout", new_page=True)
figure("01_layout_annotated.png", "Figure 1. The main window after integrating one scan.", width=5.9)
table([
    ["1.", "**Menu bar**: File, Settings, Calibration, Analysis, and Help menus (Section 12)."],
    ["2.", "**Calibration**: the calibration file (`.cal`) to use for integration."],
    ["3.", "**SPEC file**: the SPEC data file (no extension). Selecting it also fills in the user name and "
           "sets the output folder to the SPEC file's folder."],
    ["4.", "**Images**: the folder holding the Pilatus images."],
    ["5.", "**Output**: the folder where integrated `.xye` files are saved."],
    ["6.", "**User** (read from the SPEC file) and **Step** (2θ bin width, set in "
           "Settings > Integration Settings). Both are for information and cannot be edited here."],
    ["7.", "**Scan**: the scan number to integrate. Tick **Range** to enter a first and last scan "
           "instead, for in-situ or operando series."],
    ["8.", "**Integrate** starts integration of the scan or range. **Overlay** shows several patterns at "
           "once; **Contour** (with Overlay on and more than 4 patterns selected) shows them as a "
           "contour plot."],
    ["9.", "**Live Mode**: automatic integration of each new scan during an experiment, with optional "
           "event detection (Section 8)."],
    ["10.", "**Integrated Data** tab: every pattern integrated or imported this session. Click a pattern "
            "to plot it. **Events** tab: detected events (Section 9)."],
    ["11.", "**Plot area**."],
    ["12.", "**Plot toolbar**: reset view, back/forward, pan, zoom, subplot spacing, axis and line "
            "options, and save the figure as an image."],
    ["13.", "**Status bar**: progress and messages."],
], widths=[0.55, 5.95], number_col=True, size=10)
p("The dividers between the left panel and the plot, and between the controls and the data list, "
  "can be dragged to resize the panels. If the window is small, the controls scroll.")

# ------------------------------------------------------------------ 4 integrating
h("4. Integrating data")
h("4.1 Integrate one scan", 2)
steps([
    "Click **Browse** next to **Calibration** and choose the `.cal` file.",
    "Click **Browse** next to **SPEC file** and choose your SPEC file. The **User** field fills in "
    "and **Output** is set to the same folder; change Output if you want the data elsewhere.",
    "Click **Browse** next to **Images** and choose the folder with the Pilatus images.",
    "Type the scan number in **Scan** and click **Integrate**.",
])
p("A progress bar appears at the bottom of the window. When integration finishes, the pattern is saved "
  "to the output folder, added to the **Integrated Data** list, and plotted. A typical scan "
  "(16 images) takes well under a second.")
h("4.2 Integrate a range of scans", 2)
p("Tick **Range**, enter the first and last scan numbers, and click **Integrate**. The scans are "
  "integrated one after another, in order, and each is saved and added to the list as it finishes. "
  "For scans that are still being collected, use live mode instead (Section 8).")
h("4.3 Output files", 2)
p("Each scan is written as `<SPEC file name>_scan<N>.xye` in the output folder, for example "
  "`KHS10_27D_CaCO3_CuO_anneal_scan1.xye`. The file has three columns with no header:")
table([
    ["1", "2θ (degrees), bin centres at the chosen step size"],
    ["2", "Intensity: average pixel intensity in each 2θ bin, normalized to the monitor (I0) for each "
          "image and scaled to the monitor counts of the first image (a rough counts-per-pixel scale)"],
    ["3", "Uncertainty σ, from the chosen error model (Section 7)"],
], widths=[0.8, 5.7], header=["Column", "Contents"])
note("Existing files with the same name are overwritten when a scan is integrated again.")

# ------------------------------------------------------------------ 5 viewing
h("5. Viewing data")
p("Patterns in the **Integrated Data** list can be plotted in three ways:")
bullets([
    "**Single**: click a pattern in the list to plot it on its own.",
    "**Overlay**: tick **Overlay**, then click patterns to add them to (or remove them from) the "
    "plot. Each pattern gets its own colour and a legend entry.",
    "**Contour**: with **Overlay** ticked and more than 4 patterns selected, tick **Contour** to show "
    "intensity as a function of 2θ and pattern number.",
])
figure("05_overlay.png", "Figure 2. Overlay of five patterns from an in-situ run, zoomed to 10–30° "
       "with Settings > Plot Settings.", width=6.2)
h("Plot settings", 2)
p("**Settings > Plot Settings** changes line width, style and colour, the colour map used for contour "
  "and waterfall plots, the marker style, the 2θ range (Min X, Max X; **Reset X-Axis** returns to "
  "0–120°), and the intensity scale (linear, square root, or log). Click **Accept** to apply.")
figure("04_plot_settings.png", "Figure 3. Plot Settings.", width=2.1)
p("Use the plot toolbar (Figure 1, item 12) to zoom into a region, pan, go back to earlier views, or "
  "save the current plot as an image.")

# ------------------------------------------------------------------ 6 import / clear
h("6. Importing and clearing data")
bullets([
    "**File > Import Integrated Data** adds existing `.xye` files to the **Integrated Data** list. "
    "Several files can be selected at once.",
    "**File > Clear Data** removes all patterns and events from the session and clears the plot, "
    "after asking you to confirm. Your calibration file, SPEC file, image and output folders, scan "
    "numbers, and settings are kept, so you can integrate the next scan straight away. Files on "
    "disk are not touched. Live mode, if running, is stopped first.",
    "**File > Exit** closes the program.",
])

# ------------------------------------------------------------------ 7 integration settings
h("7. Integration settings")
p("**Settings > Integration Settings** controls how images are binned into a pattern. The defaults "
  "suit most measurements.")
figure("03_integration_settings.png", "Figure 4. Integration Settings.", width=2.6)
table([
    ["Min / Max 2-theta", "The 2θ range written to the output file. **Full 2-theta Range** sets "
     "0.5–180°; the file then contains only the range the detector actually covered."],
    ["Step Size", "Width of the 2θ bins, in degrees (default 0.005)."],
    ["Error Model", "How the uncertainty column σ is calculated:\n"
     "**poisson**: σ = √I. Counting statistics only.\n"
     "**azimuthal**: σ = the spread (standard deviation) of the pixel values that fall in each 2θ "
     "bin. Spotty or textured rings, where pixels at the same 2θ differ a lot, get large errors; "
     "smooth, continuous rings get errors close to counting statistics."],
    ["Lower / Upper clipping range", "Detector pixel columns outside this range (0–487) are left out, "
     "removing the detector edges. Defaults 20 and 467."],
], widths=[1.8, 4.7], header=["Setting", "Meaning"])
note("Patterns integrated with the **azimuthal** error model in versions before October 2026 have a "
     "2θ offset (about 1.6° for typical data) and a variance, not σ, in the third column. Re-integrate "
     "them with the current version. Data integrated with the **poisson** model are not affected.",
     label="Important")

# ------------------------------------------------------------------ 8 live mode
h("8. Live mode", new_page=True)
p("Live mode integrates each scan automatically as soon as it is complete and adds it to a live "
  "waterfall plot. Start it at the beginning of a run and leave it: no clicks are needed while you "
  "collect data.")
figure("06_live_annotated.png", "Figure 5. Live mode at the end of a 288-scan in-situ heating run, "
       "with detected events marked on the waterfall and listed in the Events tab.")
table([
    ["1.", "**Live integration** turns live mode on and off. **Start at**: leave blank to begin with "
           "the scan in progress (or the next scan); enter a scan number to also integrate earlier "
           "scans that are already collected."],
    ["2.", "**Detect events** (Section 9) and **Live waterfall** (redraw the waterfall after each scan; "
           "untick to look at single patterns while live mode keeps running)."],
    ["3.", "**Status**: how many scans have been integrated and which scan is expected next."],
    ["4.", "**Events** tab: events detected so far. Hover over an event for an explanation."],
    ["5.", "**Live waterfall**: intensity vs 2θ and scan number. High-confidence events are solid "
           "lines, low-confidence events dotted."],
], widths=[0.55, 5.95], number_col=True, size=10)
h("8.1 Running live mode", 2)
steps([
    "Load the calibration file, SPEC file, image folder, and output folder (Section 4.1).",
    "Leave **Start at** blank, or enter the first scan to include.",
    "Tick **Live integration**. The status line shows which scan it is waiting for.",
    "Collect data as usual. Each finished scan is integrated, saved as `.xye`, added to the data "
    "list, and drawn on the waterfall.",
    "Untick **Live integration** to stop.",
])
p("While live mode runs, the **Integrate** button is disabled.")
h("8.2 When a scan is integrated", 2)
p("A scan is integrated once it is complete:")
bullets([
    "its block in the SPEC file has all its points (for an `ascan`, the number of intervals plus "
    "one; for other scan types, when the next scan starts), and",
    "every image of the scan is present at full size, with its `.pdi` file.",
])
p("A half-written image is never integrated, and no scan is integrated twice. Scans that never "
  "produce images, such as motor alignment scans, are skipped after 30 seconds so they do not hold up "
  "the queue. The SPEC file is checked every 2 seconds.")

# ------------------------------------------------------------------ 9 events
h("9. Detected events", new_page=True)
p("An event is a scan where the diffraction pattern changed in a way that smooth thermal expansion "
  "does not explain. Detection does not know which phases are present: it flags **where** something "
  "changed, and you decide **what** changed. Events are suggestions, for example for where to start a "
  "new region in a sequential refinement; compare the patterns just before and after an event before "
  "relying on it.")
h("9.1 Reading an event", 2)
p("`scan 96 [high] area+pearson+pointwise+profile+slope, sharpening, 905 °C`")
table([
    ["scan 96", "Where the change is centred. In live mode it is reported about 6 scans later, once "
                "the detectors can confirm it."],
    ["[high] / [low]", "Confidence (below)."],
    ["area+pearson+…", "The detector families that responded (Section 9.3)."],
    ["sharpening / broadening", "Shown when peak widths changed."],
    ["905 °C", "Temperature at the event, when the furnace temperature is recorded in the SPEC file."],
], widths=[1.8, 4.7], header=["Part", "Meaning"])
h("9.2 Confidence", 2)
bullets([
    "**[high]**: two or more independent detector families agree, so the pattern changed in more than "
    "one way. Most likely a real transformation.",
    "**[low]**: only one family responded. It may be real (a subtle or gradual change that only one "
    "family sees) or a false alarm: sample movement, beam or intensity fluctuations, crowded peaks near "
    "the edge of the 2θ range, or noise early in a live run. Look at the patterns before deciding.",
])
h("9.3 Detector families", 2)
p("Each pattern is compared with the pattern 3 scans earlier, allowing peaks to shift by up to ±0.05° "
  "2θ so that thermal expansion and contraction are not events.")
table([
    ["area\n(area+ / area−)", "Integrated intensity appearing or disappearing in a 2θ region; not "
     "affected by peak width changes.",
     "area+: a phase forming or growing (reaction product, intermediate, new polymorph, "
     "crystallization). area−: a phase consumed, decomposing, melting, or becoming amorphous. Both: one "
     "phase converting into another. A loss and later a gain at the same positions: a reversible "
     "transition."],
    ["pointwise", "The same comparison, point by point on √intensity.",
     "Most sensitive to weak or minor phases. With profile but without area, usually a peak-shape "
     "change rather than a new phase."],
    ["profile", "Peak sharpness at roughly constant area.",
     "Sharpening: crystallite growth, annealing, sintering, or the finest crystallites reacting first. "
     "Broadening: strain, defects, smaller crystallites, disorder, or a peak starting to split."],
    ["slope", "How fast the whole pattern is moving away from the first scan.",
     "Gradual changes: slow reactions, steadily changing phase fractions. Also changes in heating rate."],
    ["pearson", "Abrupt change of the whole pattern between consecutive scans.",
     "Fast events; also the start or end of a hold or of cooling. Misses changes in weak phases."],
], widths=[1.15, 2.2, 3.15], header=["Family", "Measures", "May indicate"], keep=False)
h("9.4 What is not detected", 2)
bullets([
    "A phase that is present but not changing. “No events during a hold” means nothing changed beyond "
    "the thresholds, not that a reaction is complete.",
    "Changes smaller than the noise, roughly a percent of the pattern's intensity.",
    "Smooth peak shifts (thermal expansion, gradual composition change) are ignored by design. A sudden "
    "lattice change larger than 0.05° can appear as an area or pointwise event.",
])
h("9.5 Detecting events in data already loaded", 2)
p("**Analysis > Detect Events in Selected Data** runs the same detection on patterns already in the "
  "data list, whether integrated or imported. It uses the selected patterns, or all patterns if none "
  "are selected, and orders them by the scan number in their file names (`…_scanN.xye`). If the "
  "matching SPEC file is loaded, scan times and furnace temperatures (when recorded) are added to "
  "the events. At least 25 patterns are needed.")
p("Because it can use the whole series as its reference, this usually gives fewer false alarms than "
  "live mode, which only has the scans collected so far. For finished runs, it is the better choice.")
figure("07_detect_loaded.png", "Figure 6. Events detected in a loaded 437-scan run. A reversible change "
       "appears as intensity lost at scan 130 (start of the hold) and regained at scan 217 (start of "
       "cooling).")
p("**Help > Understanding Detected Events** gives this explanation inside the program, and hovering "
  "over any event in the Events tab explains that particular event.")

# ------------------------------------------------------------------ 10 calibration
h("10. Calibration")
p("The calibration file gives the direct-beam position on the detector (pixel x, y) and the "
  "sample-to-detector distance, which together set the 2θ value of every pixel. It is made with the "
  "beamline calibration script from a direct-beam scan, and its contents look like:")
for line in ("direct_beam_x  264", "direct_beam_y  79", "Sample_Detector_distance_pixels  4118.07",
             "Sample_Detector_distance_mm  708.308"):
    p(f"`{line}`", after=0)
p(after=4)
p("Use the calibration measured for your detector setup; if the detector has been moved, a new "
  "calibration is needed. **Calibration > Run Calibration** is not available in the program yet; it "
  "reminds you to use the calibration script and load the `.cal` file it writes.")

# ------------------------------------------------------------------ 11 troubleshooting
h("11. Troubleshooting")
table([
    ["“No user name was found in the SPEC file”", "The SPEC file has no `User =` line, or the wrong "
     "file was selected.", "Select the SPEC file written for your measurements."],
    ["Integration fails: image not found", "Image names must be `<user>_<SPEC file name>_scan<N>_<kkkk>.raw` "
     "(e.g. `b_stone_KHS10_27D_CaCO3_CuO_anneal_scan1_0000.raw`).",
     "Check the Images folder and that the scan number exists."],
    ["Integration fails: unexpected RAW file size", "An image is incomplete or not a Pilatus 100K "
     "`.raw` image.", "Wait for collection to finish, or check the file."],
    ["Live mode is waiting and nothing happens", "The scan in progress is not complete yet, its "
     "images are not all written, or **Start at** is beyond the last scan.",
     "Check the status line; it names the scan it is waiting for."],
    ["Many low-confidence events early in a live run", "Live detection has little history at the start.",
     "Afterwards, run Analysis > Detect Events in Selected Data on the whole run."],
    ["Peaks in the 2θ axis are off", "Wrong calibration file for this detector setup.",
     "Use the calibration measured for these scans."],
    ["Azimuthal-mode data from older versions look shifted", "Bug in versions before October 2026.",
     "Re-integrate with the current version (Section 7)."],
], widths=[1.9, 2.5, 2.1], header=["Problem", "Likely cause", "What to do"])

# ------------------------------------------------------------------ 12 menus
h("12. Menu reference")
table([
    ["File", "Import Integrated Data · Clear Data · Exit"],
    ["Settings", "Plot Settings (Section 5) · Integration Settings (Section 7)"],
    ["Calibration", "Run Calibration (not yet available; see Section 10)"],
    ["Analysis", "Detect Events in Selected Data (Section 9.5)"],
    ["Help", "Live Mode Guide · Understanding Detected Events · Open Manual PDF · About"],
], widths=[1.4, 5.1], header=["Menu", "Items"])
p("**Help > Open Manual PDF** opens `manual.pdf` from the folder the program is in. If it is not "
  "there, it opens the copy on GitHub instead, which needs access to the repository.")

doc.save(OUT)
print("saved", OUT)
